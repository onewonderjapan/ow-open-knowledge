# -*- coding: utf-8 -*-
"""構造化Markdown → 日本語ビジネスPPTX レンダラ

設計: LLMは内容(Markdown)だけ書き、体裁はコードが保証する。
  - H1 → 表紙スライド
  - H2 → 新スライド（章見出し）
  - 表  → PPTXテーブル（行数が多い場合は自動分割）
  - 箇条書き/段落 → 本文テキスト（あふれたら自動で次スライドへ）
将来: 顧客のテンプレPPTX(スライドマスタ)を --template で差し替え。

usage: python tools/md_to_pptx.py <input.md> <output.pptx> [--subtitle "..."]
"""
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

NAVY = RGBColor(0x1A, 0x2B, 0x4A)
ACCENT = RGBColor(0x2F, 0x6F, 0xED)
GREY = RGBColor(0x66, 0x70, 0x80)
FONT = "Yu Gothic UI"

SLIDE_W, SLIDE_H = Cm(33.87), Cm(19.05)  # 16:9
MAX_TABLE_ROWS = 9      # 1スライドに載せる表の最大行(ヘッダ除く)
MAX_BODY_LINES = 13     # 1スライドに載せる本文行数


def parse_md(src: str) -> list[dict]:
    """Markdown → ブロック列 [{type: title|section|table|lines, ...}]"""
    blocks, lines_buf, table_buf = [], [], []

    def flush_lines():
        nonlocal lines_buf
        if lines_buf:
            blocks.append({"type": "lines", "lines": lines_buf})
            lines_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            blocks.append({"type": "table", "rows": table_buf})
            table_buf = []

    for raw in src.splitlines():
        l = raw.rstrip()
        if l.startswith("|"):
            flush_lines()
            cells = [c.strip() for c in l.strip("|").split("|")]
            if not all(re.fullmatch(r"[-: ]*", c) for c in cells):
                table_buf.append(cells)
            continue
        flush_table()
        if l.startswith("# "):
            flush_lines()
            blocks.append({"type": "title", "text": l[2:].strip()})
        elif l.startswith("## "):
            flush_lines()
            blocks.append({"type": "section", "text": l[3:].strip()})
        elif l.strip():
            lines_buf.append(l)
    flush_lines()
    flush_table()
    return blocks


def _clean(s: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", s)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.current_section = ""

    def _slide(self):
        return self.prs.slides.add_slide(self.blank)

    def title_slide(self, title: str, subtitle: str):
        s = self._slide()
        bar = s.shapes.add_textbox(Cm(0), Cm(7.2), SLIDE_W, Cm(0.1))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
        tb = s.shapes.add_textbox(Cm(2.5), Cm(7.6), Cm(28.8), Cm(3))
        p = tb.text_frame.paragraphs[0]
        r = p.add_run(); r.text = title
        r.font.size, r.font.bold, r.font.name = Pt(32), True, FONT
        r.font.color.rgb = NAVY
        st = s.shapes.add_textbox(Cm(2.5), Cm(11), Cm(28.8), Cm(1.5))
        p2 = st.text_frame.paragraphs[0]
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.size, r2.font.name, r2.font.color.rgb = Pt(14), FONT, GREY

    def section_header(self, s, text: str):
        tb = s.shapes.add_textbox(Cm(1.2), Cm(0.7), Cm(31.4), Cm(1.6))
        p = tb.text_frame.paragraphs[0]
        r = p.add_run(); r.text = text
        r.font.size, r.font.bold, r.font.name = Pt(20), True, FONT
        r.font.color.rgb = NAVY
        ln = s.shapes.add_textbox(Cm(1.2), Cm(2.35), Cm(31.4), Cm(0.06))
        ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()

    def body_slide(self, lines: list[str], cont: bool = False):
        s = self._slide()
        self.section_header(s, self.current_section + ("（続き）" if cont else ""))
        tb = s.shapes.add_textbox(Cm(1.6), Cm(3.0), Cm(30.6), Cm(15.0))
        tf = tb.text_frame; tf.word_wrap = True
        first = True
        for l in lines:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            m = re.match(r"^(\s*)[-*・]\s*(.+)$", l)
            if m:
                p.level = min(2, len(m.group(1)) // 2)
                text = "・" + _clean(m.group(2))
            else:
                text = _clean(l)
            r = p.add_run(); r.text = text
            r.font.size, r.font.name = Pt(14), FONT
            p.space_after = Pt(6)
        return s

    def table_slides(self, rows: list[list[str]]):
        header, body = rows[0], rows[1:]
        chunks = [body[i:i + MAX_TABLE_ROWS] for i in range(0, len(body), MAX_TABLE_ROWS)] or [[]]
        for ci, chunk in enumerate(chunks):
            s = self._slide()
            self.section_header(s, self.current_section + ("（続き）" if ci else ""))
            ncols, nrows = len(header), len(chunk) + 1
            tbl = s.shapes.add_table(nrows, ncols, Cm(1.2), Cm(3.0), Cm(31.4), Cm(1.0) * nrows).table
            for j, h in enumerate(header):
                c = tbl.cell(0, j); c.text = _clean(h)
                c.fill.solid(); c.fill.fore_color.rgb = NAVY
                for p in c.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for r in p.runs:
                        r.font.size, r.font.bold, r.font.name = Pt(11), True, FONT
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            for i, row in enumerate(chunk, 1):
                for j in range(ncols):
                    c = tbl.cell(i, j); c.text = _clean(row[j]) if j < len(row) else ""
                    for p in c.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size, r.font.name = Pt(10.5), FONT

    def render(self, blocks: list[dict], subtitle: str):
        pend: list[str] = []

        def flush_pending(cont=False):
            nonlocal pend
            while pend:
                chunk, pend = pend[:MAX_BODY_LINES], pend[MAX_BODY_LINES:]
                self.body_slide(chunk, cont=cont)
                cont = True

        for b in blocks:
            if b["type"] == "title":
                self.title_slide(b["text"], subtitle)
            elif b["type"] == "section":
                flush_pending()
                self.current_section = b["text"]
            elif b["type"] == "lines":
                pend.extend(b["lines"])
            elif b["type"] == "table":
                flush_pending()
                self.table_slides(b["rows"])
        flush_pending()


def main():
    src, dst = sys.argv[1], sys.argv[2]
    subtitle = sys.argv[sys.argv.index("--subtitle") + 1] if "--subtitle" in sys.argv else ""
    blocks = parse_md(Path(src).read_text(encoding="utf-8"))
    deck = Deck()
    deck.render(blocks, subtitle)
    deck.prs.save(dst)
    n = len(deck.prs.slides._sldIdLst)
    print(f"saved {n} slides -> {dst}")


if __name__ == "__main__":
    main()
