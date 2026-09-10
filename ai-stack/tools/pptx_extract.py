# -*- coding: utf-8 -*-
"""PPTX → テキスト抽出（スライド順、表対応、ノート含む）

usage: python tools/pptx_extract.py <input.pptx> <output.md>
"""
import sys
from pathlib import Path

from pptx import Presentation


def shape_text(shape) -> list[str]:
    out = []
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if t:
                out.append(("  " * para.level) + ("- " if para.level else "") + t)
    if shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.replace("\n", " ").strip() for c in row.cells]
            out.append("| " + " | ".join(cells) + " |")
    if shape.shape_type == 6:  # GROUP
        for s in shape.shapes:
            out.extend(shape_text(s))
    return out


def extract(path: str) -> str:
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            lines.extend(shape_text(shape))
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                lines.append(f"(ノート: {note})")
        parts.append(f"## スライド{i}\n" + "\n".join(lines))
    return "\n\n".join(parts)


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    text = extract(src)
    Path(dst).write_text(text, encoding="utf-8")
    print(f"extracted {len(text)} chars -> {dst}")
